from __future__ import annotations

from pathlib import Path

from Tool.contracts.canonical import CanonicalDocument, DocumentMeta, Fragment, Section, TableData
from Tool.normalizers import detect_doc_type, extract_terms, normalize_text


def parse_xlsx(file_path: Path, manifest: dict) -> CanonicalDocument:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to parse XLSX files.") from exc

    workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    sections: list[Section] = []
    fragments: list[Fragment] = []
    tables: list[TableData] = []

    for sheet_index, sheet_name in enumerate(workbook.sheetnames, start=1):
        sheet = workbook[sheet_name]
        section_id = f"sec-{sheet_index}"
        sections.append(Section(section_id=section_id, title=sheet_name, level=1, page_range=[sheet_index, sheet_index]))

        rows: list[list[str]] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [normalize_text("" if value is None else str(value)) for value in row]
            if not any(cells):
                continue
            rows.append(cells)
            fragments.append(
                Fragment(
                    fragment_id=f"frag-{sheet_index}-{row_index}",
                    section_id=section_id,
                    fragment_type="row",
                    text=" | ".join(cells),
                    anchors={"sheet": sheet_name, "row_index": row_index},
                )
            )

        if rows:
            tables.append(
                TableData(
                    table_id=f"tbl-{sheet_index}",
                    section_id=section_id,
                    page=sheet_index,
                    rows=rows,
                    anchors={"sheet": sheet_name},
                )
            )

    title = workbook.sheetnames[0] if workbook.sheetnames else file_path.stem
    meta = DocumentMeta(
        document_id=manifest["document_id"],
        title=title,
        source_path=manifest["stored_path"],
        file_name=file_path.name,
        source_type="xlsx",
        doc_type=detect_doc_type(title, file_path.name),
        checksum=manifest.get("checksum", ""),
        metadata={"manifest_path": manifest.get("manifest_path", ""), "sheet_count": len(workbook.sheetnames)},
    )

    return CanonicalDocument(
        document=meta,
        sections=sections,
        fragments=fragments,
        tables=tables,
        figures=[],
        terms=extract_terms([fragment.text for fragment in fragments]),
        entities=[],
        parse_status="parsed" if fragments else "failed",
        source_anchors=[
            {"fragment_id": fragment.fragment_id, "anchors": fragment.anchors}
            for fragment in fragments
        ],
        errors=[] if fragments else ["No rows extracted from XLSX."],
    )
